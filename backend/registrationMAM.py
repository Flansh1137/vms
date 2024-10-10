from flask import Flask, request, jsonify
from flask_cors import CORS
import json
import os
import base64
from tensorflow.keras import backend as K 
from livePredict_flask import process_videos_sequentially, datewise_filter, test
from newDataset import Train_model
from openpyxl import load_workbook
import pandas as pd
import traceback
import cv2
import mysql.connector
#from sqlDB import mydb
import pickle
from datetime import datetime
#from pytube import YouTube
cam= None


app1 = Flask(__name__)
CORS(app1)

DATA_FILE = 'NewRegistration.xlsx'
DATA_FORM_FILE = 'DataFormAnalysis.json'
IMAGE_DIR = 'images'

# Initialize registration data file
if not os.path.exists(DATA_FILE):
    with open(DATA_FILE, 'w') as file:
        json.dump([], file)

# Initialize data form file
if not os.path.exists(DATA_FORM_FILE):
    with open(DATA_FORM_FILE, 'w') as file:
        json.dump([], file)

# Ensure image directory exists
if not os.path.exists(IMAGE_DIR):
    os.makedirs(IMAGE_DIR)

def getconnection():

    return mysql.connector.connect(
        host="xbit.c3i4g8kgexgt.us-east-1.rds.amazonaws.com",
        user="admin",
        password="December2020",
        database="myDB"
    )
mydb=getconnection()

@app1.route('/')
def index():
    return jsonify({'message': 'Welcome to the API!'})


@app1.route('/main-program', methods=['POST'])
def main_program():
       
        print("mainnnnn")
        cam=cv2.VideoCapture("videos/modi.mp4")
        test(cam)
          
        return jsonify({'message': 'Live Recording on'})

@app1.route('/new-registration', methods=['POST'])
def new_registration():
    try:
        data = request.json
        print('Received data:', data)
          # Log the received data to console
        id=data.get('id')
        name=data.get('name')
        phone=data.get("phone")
        age=data.get("age")
        bloodGroup=data.get("bloodGroup")
        height=data.get("height")
        weight=data.get("weight")
        address=data.get("address")
        emergencyContact=data.get("emergencyContact")
        print(name)
        data=[[id, name, phone, age, bloodGroup, height, weight, address, emergencyContact]]
        wb = load_workbook(DATA_FILE)
        ws=wb.active
        for row in data:
            ws.append(row)
        wb.save(DATA_FILE)

        mycursor=mydb.cursor()
        sql="INSERT INTO Members(id, member_name, phone,age, bloodgroup, height, weight, address, emergencyContact) VALUES(%s, %s, %s,%s, %s, %s,%s, %s, %s)"
        val=(id, name, age, phone, bloodGroup, height, weight,address, emergencyContact)
        mycursor.execute(sql, val)
        mydb.commit()

        return jsonify({'message': 'Registration successful', 'data': data}), 200
    except Exception as e:
        print('Error processing request:', e)
        return jsonify({'error': 'Internal Server Error'}), 500

@app1.route('/upload-image', methods=['POST'])
def upload_image():
    try:
        data = request.json
        user_id = data.get('id')
        user_name = data.get('name')
        images = data.get('images', [])
        saved_images = []
        simages=[]

        # Check if user_id and user_name are not None
        if not user_id or not user_name:
            return jsonify({'error': 'ID and Name are required.'}), 400

        user_dir = os.path.join(IMAGE_DIR, f'{user_name}_{user_id}')
        if not os.path.exists(user_dir):
            os.makedirs(user_dir)

        for idx, image_data in enumerate(images):
            image_base64 = image_data.split(',')[1]
            image_bytes = base64.b64decode(image_base64)
            image_path = os.path.join(user_dir, f'image_{idx + 1}.jpeg')
            simages.append(image_bytes)
            serialized_data = pickle.dumps(simages)
            with open(image_path, 'wb') as image_file:
                image_file.write(image_bytes)
            saved_images.append(image_path)

        mycursor=mydb.cursor()

        sql = "UPDATE Members SET images = %s WHERE id = %s AND member_name = %s"
        val = (serialized_data, user_id, user_name)
        mycursor.execute(sql, val)
        mydb.commit()

        path="images/"
       
        Train_model(path)

        return jsonify({'message': 'Images uploaded successfully', 'images': saved_images}), 200
    except Exception as e:
        print('Error processing request:', e)
        return jsonify({'error': 'Internal Server Error'}), 500
    


@app1.route('/stop-camera', methods=["GET"])
def stop_camera():
    global cam
    if cam is None:
        return jsonify({'message': 'Camera is not open'}), 400
    with open("camera_status.txt", 'w') as file:
        file.write('stop')
    cam.release()
    cam=None

    return jsonify({'message': 'camera stopped successfully'}), 200



@app1.route('/unknown-data-analysis', methods=['GET'])
def unknown_data_analysis():
    date_unique_count_dict={}
    try:
        global mydb
        # Get start and end dates from query parameters
        start_date_str = request.args.get('startDate')
        end_date_str = request.args.get('endDate')
        
        if not start_date_str or not end_date_str:
            raise ValueError("Missing 'startDate' or 'endDate' parameter")
       
        if not mydb.is_connected():
            mydb = getconnection()
        mycursor = mydb.cursor()
        
        # Prepare your SQL statement with placeholders
        try:
            sql = 'SELECT DATE(DATE_TIME) AS date, COUNT(DISTINCT UNKNOWN_NAME) AS unique_count FROM UNKNOWN_DB WHERE DATE(DATE_TIME) BETWEEN %s AND %s GROUP BY DATE(DATE_TIME)'
            mycursor.execute(sql, (start_date_str, end_date_str))
            results = mycursor.fetchall()
            print("unknownnnnnnn ", results)
            
            date_unique_count_dict = {date.strftime('%Y-%m-%d'): count for date, count in results}
            #date_unique_count_dict = {row[0].strftime('%Y-%m-%d'): row[1] for row in results}
            
        except Exception as e:
            print("An error occurred:", e)       
        
    except Exception as e:
        # Log the error with traceback
        traceback_str = traceback.format_exc()
        print(f"Error: {e}\nTraceback: {traceback_str}")
        return jsonify({"error": str(e)}), 500
    return jsonify(date_unique_count_dict)



@app1.route('/data-analysis', methods=['GET'])
def data_analysis():
    date_unique_count_dict={}
    try:
        global mydb
        # Get start and end dates from query parameters
        start_date_str = request.args.get('startDate')
        end_date_str = request.args.get('endDate')
        
        if not start_date_str or not end_date_str:
            raise ValueError("Missing 'startDate' or 'endDate' parameter")
       
        if not mydb.is_connected():
            mydb = getconnection()
        mycursor = mydb.cursor()
        
        # Prepare your SQL statement with placeholders
        try:
            sql = 'SELECT DATE(DATE_TIME) AS date, COUNT(DISTINCT MEM_NAME) AS unique_count FROM LIVEVIDEODATA WHERE DATE(DATE_TIME) BETWEEN %s AND %s GROUP BY DATE(DATE_TIME)'
            mycursor.execute(sql, (start_date_str, end_date_str))
            results = mycursor.fetchall()
        
            
            date_unique_count_dict = {date.strftime('%Y-%m-%d'): count for date, count in results}
            #date_unique_count_dict = {row[0].strftime('%Y-%m-%d'): row[1] for row in results}
            
        except Exception as e:
            print("An error occurred:", e)       
        
    except Exception as e:
        # Log the error with traceback
        traceback_str = traceback.format_exc()
        print(f"Error: {e}\nTraceback: {traceback_str}")
        return jsonify({"error": str(e)}), 500
    return jsonify(date_unique_count_dict)





@app1.route('/names-by-date', methods=['GET'])
def names_by_date():
    global mydb
    try:
        # Get the date from query parameters
        date_str = request.args.get('date')
        
        if not date_str:
            raise ValueError("Missing 'date' parameter")

        # Convert date to datetime object
        date = pd.to_datetime(date_str, format='%Y-%m-%d').date()  # Use .date() to get a date object

        # Ensure database connection
        if not mydb or not mydb.is_connected():
            mydb = getconnection()
        
        if not mydb:
            return jsonify({"error": "Database connection failed"}), 500

        mycursor = mydb.cursor()

        # SQL query to fetch distinct names for the specific date
        sql = 'SELECT DISTINCT MEM_NAME FROM LIVEVIDEODATA WHERE DATE(DATE_TIME) = %s'
    
        # Debugging: Print the SQL and parameters
        print("Executing SQL:", sql)
        print("With parameter:", date)

        # Execute the SQL query
        mycursor.execute(sql, (date,))  # Ensure to pass the date as a tuple

        results = mycursor.fetchall()

        # Debugging: Print the results
        print("names==", results)

        # Extract names from the results
        names = [row[0] for row in results]  # Assuming MEM_NAME is the first column

        return jsonify(names)

    except Exception as e:
        # Log the error with traceback
        traceback_str = traceback.format_exc()
        print(f"Error: {e}\nTraceback: {traceback_str}")
        return jsonify({"error": str(e)}), 500




app1.run(port=5000, debug=True)  # Ensure backend runs on port 5000
